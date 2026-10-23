"""Generate the binary sample documents (DOCX glossary, PDF reporting policy).

The Markdown docs in sample_data/docs/ are hand-written; these two exist so the
DOCX and PDF parsers are exercised by the same evaluation. Re-run after editing
the text below:

    python scripts/generate_sample_docs.py
"""
from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "sample_data" / "docs"

GLOSSARY = [
    ("Customer", "A person or organisation with at least one recorded purchase or an active "
     "subscription. In retail data a customer is identified by the loyalty-card code; in sales "
     "data by the account code."),
    ("Region", "One of the five sales territories (North, South, East, West, Central). Each region "
     "has a regional manager who signs off its monthly forecast."),
    ("Sales representative", "The account owner credited with an order. Self-service orders have no "
     "sales representative and are not counted in individual targets."),
    ("Channel", "The route to market. Online means the company web store; Retail means company-owned "
     "stores; Partner means resellers on a revenue-share contract; Distributor means wholesale."),
    ("Payment method", "How a retail transaction was paid: Cash, Card, UPI, Wallet or NetBanking. "
     "UPI is India's instant real-time payment system and is the fastest-growing method."),
    ("Rating", "The post-purchase satisfaction score from 1 (very poor) to 5 (excellent). Surveys are "
     "sent by SMS within one hour of purchase; unanswered surveys are recorded as a 3."),
    ("Conversion rate", "The percentage of website sessions that end in a completed purchase."),
    ("Bounce rate", "The percentage of website sessions that viewed only one page before leaving."),
    ("Churn", "A subscription customer who cancels and does not resubscribe within 30 days."),
]

POLICY_PAGES = [
    ("Reporting Policy", [
        "This policy governs how business metrics are calculated and reported.",
        "Fiscal year. The fiscal year runs from 1 April to 31 March. Q1 is April to June, Q2 is July "
        "to September, Q3 is October to December and Q4 is January to March.",
        "Currency. Sales figures are reported in US dollars. Retail store figures are reported in "
        "Indian rupees and are never converted in operational dashboards.",
        "Data refresh. Sales data refreshes nightly at 02:00 UTC. Retail data refreshes every hour. "
        "Web analytics refresh daily at 06:00 UTC.",
    ]),
    ("Targets and Thresholds", [
        "Satisfaction target. At least 60 percent of retail transactions should be satisfied "
        "purchases (rated 4 or 5).",
        "Web targets. The target conversion rate is 3.5 percent. The target bounce rate is below "
        "55 percent. Cost per session should stay below 6 US dollars.",
        "Margin review. Any product category whose margin rate falls below 30 percent in a quarter "
        "is escalated to the pricing committee.",
        "Refunds. Refunds are recorded as separate negative transactions and are excluded from "
        "revenue reporting until the refund is approved.",
    ]),
    ("Data Privacy", [
        "Customer emails and phone numbers are personal data. They must never be exported, "
        "shared outside the analytics team, or included in reports.",
        "Access to customer master data is reviewed every quarter by the data protection officer.",
    ]),
]


def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Business Glossary", level=0)
    document.add_paragraph("Shared vocabulary used by sales, retail and marketing teams.")
    for term, definition in GLOSSARY:
        document.add_heading(term, level=1)
        document.add_paragraph(definition)
    document.save(path)


def build_pdf(path: Path) -> None:
    import textwrap

    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams["pdf.fonttype"] = 42  # embed TrueType so text is extractable
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    with PdfPages(path) as pdf:
        for title, paragraphs in POLICY_PAGES:
            fig = plt.figure(figsize=(8.27, 11.69))  # A4
            fig.text(0.08, 0.94, title, fontsize=18, weight="bold", va="top")
            y = 0.89
            for para in paragraphs:
                for line in textwrap.wrap(para, 85):
                    fig.text(0.08, y, line, fontsize=10.5, va="top")
                    y -= 0.022
                y -= 0.014
            pdf.savefig(fig)
            plt.close(fig)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / "business_glossary.docx")
    build_pdf(OUT / "reporting_policy.pdf")
    print(f"wrote {OUT / 'business_glossary.docx'} and {OUT / 'reporting_policy.pdf'}")
